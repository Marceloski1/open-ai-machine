"""Fixtures binarias reales para probar la conversion de cada formato.

Se generan en tiempo de test en vez de versionar binarios: el repositorio se
mantiene libre de blobs y cada fixture es reproducible. El .docx reutiliza el
generador de plantillas del propio repositorio.
"""

from __future__ import annotations

import importlib.util
import struct
import zlib
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
BUILD_TEMPLATE_PATH = REPO_ROOT / "tools" / "build-template" / "build-template.py"


def _build_pdf(lines: list[str]) -> bytes:
    """PDF 1.4 minimo con texto extraible y tabla xref coherente."""
    content = "BT /F1 24 Tf 72 720 Td 28 TL\n"
    for line in lines:
        content += f"({line}) Tj T*\n"
    content += "ET"
    stream = content.encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _build_png(width: int = 4, height: int = 4) -> bytes:
    """PNG opaco minimo, sin EXIF ni metadatos."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    path = tmp_path / "informe.pdf"
    path.write_bytes(_build_pdf([
        "Informe de descubrimiento",
        "Fase 0 del pipeline machine.",
    ]))
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    spec = importlib.util.spec_from_file_location("build_template", BUILD_TEMPLATE_PATH)
    build_template = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(build_template)
    return build_template.build(tmp_path / "acta.docx")


@pytest.fixture
def xlsx_file(tmp_path: Path) -> Path:
    import xlsxwriter

    path = tmp_path / "presupuesto.xlsx"
    workbook = xlsxwriter.Workbook(str(path))
    worksheet = workbook.add_worksheet("Fase0")
    rows = [["concepto", "monto"], ["descubrimiento", 1500], ["propuesta", 2300]]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            worksheet.write(row_index, column_index, value)
    workbook.close()
    return path


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    from pptx import Presentation

    path = tmp_path / "pitch.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Propuesta de negocio"
    slide.placeholders[1].text = "Mercado, problema y solucion"
    presentation.save(str(path))
    return path


@pytest.fixture
def png_file(tmp_path: Path) -> Path:
    path = tmp_path / "pizarra.png"
    path.write_bytes(_build_png())
    return path
